import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.ids import next_req_id
from app.models import CaseFolder, Project, Requirement, User
from app.services.excel_import import _cell, auto_map, read_tabular

REQ_FIELDS = ["reqid", "title", "description", "priority", "module"]

REQ_ALIASES = {
    "requirement id": "reqid",
    "req id": "reqid",
    "reqid": "reqid",
    "id": "reqid",
    "requirement": "title",
    "title": "title",
    "name": "title",
    "summary": "title",
    "description": "description",
    "desc": "description",
    "priority": "priority",
    "module": "module",
    "module name": "module",
    "folder": "module",
}


def req_auto_map(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for header in headers:
        key = re.sub(r"[^a-z0-9]+", "", (header or "").strip().lower())
        spaced = re.sub(r"\s+", " ", (header or "").strip().lower())
        target = REQ_ALIASES.get(spaced) or (key if key in REQ_FIELDS else None)
        if target:
            mapping[header] = target
    return mapping


def build_req_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    headers = ["RequirementID", "Title", "Description", "Priority", "Module"]
    fill = PatternFill("solid", fgColor="2563EB")
    font = Font(color="FFFFFF", bold=True)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 24
    ws.append(["", "User can log in with valid credentials", "Login with email and password", "Medium", "Feature"])
    help_sheet = wb.create_sheet("Help")
    help_sheet["A1"] = "RequirementID is optional. Empty IDs get REQ-{PROJECT}-{nnn}."
    help_sheet["A2"] = "Title is required. Duplicate RequirementIDs in the file or project are skipped."
    help_sheet["A3"] = "Priority: High / Medium / Low. Module must already exist as a sprint module; it is not created."
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm_priority(value: str) -> str:
    raw = (value or "MEDIUM").strip().upper()
    aliases = {"MED": "MEDIUM", "MID": "MEDIUM"}
    raw = aliases.get(raw, raw)
    return raw if raw in {"HIGH", "MEDIUM", "LOW"} else "MEDIUM"


def classify_requirement_rows(
    rows: list[dict],
    mapping: dict[str, str],
    existing_ids: set[str],
) -> dict:
    items = []
    seen: dict[str, int] = {}
    duplicate_in_file: list[str] = []
    duplicate_existing: list[str] = []
    issues: list[dict] = []
    for index, row in enumerate(rows, 2):
        req_id = _cell(row, mapping, "reqid").strip()
        title = _cell(row, mapping, "title").strip()
        description = _cell(row, mapping, "description").strip()
        priority = _norm_priority(_cell(row, mapping, "priority"))
        module = _cell(row, mapping, "module").strip()
        item = {
            "row": index,
            "req_id": req_id,
            "title": title,
            "description": description,
            "priority": priority,
            "module": module,
        }
        if not title:
            item["status"] = "error"
            item["issue"] = "Title is required"
            issues.append({"row": index, "req_id": req_id, "message": "Title is required"})
        elif req_id:
            key = req_id.casefold()
            if key in seen:
                item["status"] = "duplicate_file"
                item["issue"] = f"Duplicate ID in file (also row {seen[key]})"
                duplicate_in_file.append(req_id)
                issues.append({"row": index, "req_id": req_id, "message": "Duplicate Requirement ID in this file"})
            elif key in {value.casefold() for value in existing_ids}:
                item["status"] = "duplicate_existing"
                item["issue"] = "Requirement ID already exists"
                duplicate_existing.append(req_id)
                issues.append({"row": index, "req_id": req_id, "message": "Requirement ID already exists in this project"})
            else:
                item["status"] = "ok"
                seen[key] = index
        else:
            item["status"] = "ok"
        items.append(item)
    return {
        "items": items,
        "duplicate_in_file": sorted(set(duplicate_in_file)),
        "duplicate_existing": sorted(set(duplicate_existing)),
        "issues": issues,
        "ok_count": sum(1 for item in items if item.get("status") == "ok"),
    }


def folder_by_name(db: Session, project_id: str, name: str) -> str | None:
    if not name:
        return None
    wanted = name.casefold()
    match = (
        db.query(CaseFolder)
        .filter(CaseFolder.project_id == project_id, CaseFolder.kind == "WORK")
        .all()
    )
    for folder in match:
        if folder.name.casefold() == wanted:
            return folder.id
    return None


def commit_requirements(db: Session, project: Project, _user: User, classified: dict) -> dict:
    created = 0
    skipped = 0
    existing = {
        row.req_id.casefold()
        for row in db.query(Requirement).filter(Requirement.project_id == project.id).all()
    }
    for item in classified.get("items") or []:
        if item.get("status") != "ok":
            skipped += 1
            continue
        req_id = (item.get("req_id") or "").strip() or next_req_id(db, project)
        if req_id.casefold() in existing:
            skipped += 1
            continue
        req = Requirement(
            project_id=project.id,
            req_id=req_id,
            title=item["title"],
            description=item.get("description") or "",
            priority=item.get("priority") or "MEDIUM",
            folder_id=folder_by_name(db, project.id, item.get("module") or ""),
        )
        db.add(req)
        existing.add(req_id.casefold())
        created += 1
    db.commit()
    return {"created": created, "skipped": skipped}


def preview_requirements(filename: str, data: bytes, existing_ids: set[str]) -> dict:
    headers, rows = read_tabular(filename, data)
    mapping = req_auto_map(headers) if headers else {}
    classified = classify_requirement_rows(rows, mapping, existing_ids)
    return {
        "headers": headers,
        "mapping": mapping,
        "row_count": len(rows),
        "rows": rows,
        "preview_rows": rows[:25],
        "fields": REQ_FIELDS,
        **classified,
    }
