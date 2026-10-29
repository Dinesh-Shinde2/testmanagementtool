from collections import defaultdict
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.models import (
    CASE_STATUSES,
    CASE_TYPES,
    EXEC_STATUSES,
    PRIORITIES,
    CaseFolder,
    Defect,
    Execution,
    ExecutionEvent,
    Module,
    Project,
    Requirement,
    RequirementTestCase,
    Sprint,
    TestCase,
    TestPlan,
    TestRun,
    User,
)
from app.serialize import iso
from app.services.execute import defect_is_open, defect_stage, defect_story, empty_flow
from app.services.folders import sprint_cases, sprint_work_folder_ids


def _counts(statuses: list[str]) -> dict:
    keys = ("PASS", "FAIL", "BLOCKED", "SKIPPED", "NOT_EXECUTED")
    data = {k.lower(): statuses.count(k) for k in keys}
    data["total"] = len(statuses)
    executed = data["pass"] + data["fail"] + data["blocked"] + data["skipped"]
    data["executed"] = executed
    data["pass_rate"] = round((data["pass"] / executed) * 100, 1) if executed else 0.0
    return data


def latest_run_for_sprint(db: Session, sprint_id: str) -> TestRun | None:
    return (
        db.query(TestRun)
        .filter(TestRun.sprint_id == sprint_id)
        .order_by(TestRun.started_at.desc())
        .first()
    )


def parse_day(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid date. Use YYYY-MM-DD.") from exc


def _in_date_range(executed_at: datetime | None, date_from: date | None, date_to: date | None) -> bool:
    if not date_from and not date_to:
        return True
    if executed_at is None:
        return False
    day = executed_at.date()
    if date_from and day < date_from:
        return False
    if date_to and day > date_to:
        return False
    return True


def _executions_by_case(rows: list[Execution]) -> dict[str, Execution]:
    """One row per test case: latest applicable execution (executed_at, else latest in list)."""
    by_case: dict[str, Execution] = {}
    for row in rows:
        current = by_case.get(row.test_case_id)
        if current is None:
            by_case[row.test_case_id] = row
            continue
        left = current.executed_at or datetime.min
        right = row.executed_at or datetime.min
        if right >= left:
            by_case[row.test_case_id] = row
    return by_case


def _current_executions_by_case(rows: list[Execution]) -> dict[str, Execution]:
    """Latest current result per case, preferring an actual result over a later Not Run row."""
    by_case: dict[str, Execution] = {}
    for row in rows:
        current = by_case.get(row.test_case_id)
        if current is None:
            by_case[row.test_case_id] = row
            continue
        left_done = current.status != "NOT_EXECUTED"
        right_done = row.status != "NOT_EXECUTED"
        if right_done and not left_done:
            by_case[row.test_case_id] = row
            continue
        if left_done and not right_done:
            continue
        left = current.executed_at or datetime.min
        right = row.executed_at or datetime.min
        if right >= left:
            by_case[row.test_case_id] = row
    return by_case


def _latest_status(execution: Execution | None, date_from: date | None, date_to: date | None) -> str:
    if not execution:
        return "NOT_EXECUTED"
    if execution.status == "NOT_EXECUTED":
        return "NOT_EXECUTED"
    if not _in_date_range(execution.executed_at, date_from, date_to):
        return "NOT_EXECUTED"
    return execution.status


def _kpis_for_cases(
    cases: list[TestCase],
    exec_by_case: dict[str, Execution],
    date_from: date | None,
    date_to: date | None,
) -> dict:
    statuses = [_latest_status(exec_by_case.get(case.id), date_from, date_to) for case in cases]
    return _counts(statuses)


def sprint_metrics_map(db: Session, project_id: str, sprints: list[Sprint]) -> dict[str, dict]:
    """Unique-case latest-result metrics per sprint, matching Dashboard / Execute / RTM."""
    empty = {
        "module_count": 0,
        "total": 0,
        "executed": 0,
        "pass": 0,
        "fail": 0,
        "blocked": 0,
        "skipped": 0,
        "not_executed": 0,
        "pass_rate": 0.0,
    }
    result = {sprint.id: dict(empty) for sprint in sprints}
    if not sprints:
        return result
    sprint_ids = [sprint.id for sprint in sprints]
    folders = (
        db.query(CaseFolder)
        .filter(CaseFolder.sprint_id.in_(sprint_ids), CaseFolder.kind == "WORK")
        .all()
    )
    folder_sprint = {folder.id: folder.sprint_id for folder in folders}
    for folder in folders:
        result[folder.sprint_id]["module_count"] += 1
    folder_ids = list(folder_sprint)
    cases_by_sprint: dict[str, list[TestCase]] = defaultdict(list)
    if folder_ids:
        cases = (
            db.query(TestCase)
            .filter(
                TestCase.project_id == project_id,
                TestCase.folder_id.in_(folder_ids),
                TestCase.status != "DEPRECATED",
            )
            .all()
        )
        for case in cases:
            sprint_id = folder_sprint.get(case.folder_id)
            if sprint_id:
                cases_by_sprint[sprint_id].append(case)
    runs = db.query(TestRun).filter(TestRun.sprint_id.in_(sprint_ids)).all()
    run_ids = [run.id for run in runs]
    execs_by_sprint: dict[str, list[Execution]] = defaultdict(list)
    run_sprint = {run.id: run.sprint_id for run in runs}
    if run_ids:
        for row in db.query(Execution).filter(Execution.run_id.in_(run_ids)).all():
            sid = run_sprint.get(row.run_id)
            if sid:
                execs_by_sprint[sid].append(row)
    for sprint in sprints:
        case_rows = cases_by_sprint.get(sprint.id, [])
        by_case = _current_executions_by_case(execs_by_sprint.get(sprint.id, []))
        stats = result[sprint.id]
        stats.update(_kpis_for_cases(case_rows, by_case, None, None))
    return result


def sprint_links_map(db: Session, sprint_ids: list[str]) -> dict[str, dict]:
    empty = {"modules": 0, "cases": 0, "plans": 0, "runs": 0, "requirements": 0}
    result = {sprint_id: dict(empty) for sprint_id in sprint_ids}
    if not sprint_ids:
        return result
    folders = (
        db.query(CaseFolder)
        .filter(CaseFolder.sprint_id.in_(sprint_ids), CaseFolder.kind == "WORK")
        .all()
    )
    folder_sprint = {folder.id: folder.sprint_id for folder in folders}
    for folder in folders:
        result[folder.sprint_id]["modules"] += 1
    folder_ids = list(folder_sprint)
    if folder_ids:
        cases = (
            db.query(TestCase)
            .filter(TestCase.folder_id.in_(folder_ids), TestCase.status != "DEPRECATED")
            .all()
        )
        for case in cases:
            sprint_id = folder_sprint.get(case.folder_id)
            if sprint_id:
                result[sprint_id]["cases"] += 1
    for plan in db.query(TestPlan).filter(TestPlan.sprint_id.in_(sprint_ids)).all():
        result[plan.sprint_id]["plans"] += 1
    for run in db.query(TestRun).filter(TestRun.sprint_id.in_(sprint_ids)).all():
        result[run.sprint_id]["runs"] += 1
    for req in db.query(Requirement).filter(Requirement.sprint_id.in_(sprint_ids)).all():
        result[req.sprint_id]["requirements"] += 1
    return result


def _pct_of(part: int, total: int) -> float | None:
    if not total:
        return None
    return round((part / total) * 100, 1)


def case_history_stats(
    db: Session,
    project_id: str,
    cases: list[TestCase],
    sprint_ids: list[str],
    date_from: date | None,
    date_to: date | None,
    current_by_case: dict[str, Execution],
) -> dict[str, dict]:
    """Per-case fail/retest history from events, executions and bug lifecycle.

    Passed After Retest requires a recorded FAIL followed by a later PASS
    (execution event or closed-bug retest). Current PASS alone is not enough.
    """
    blank = {"fail_count": 0, "retest_pass_count": 0, "ever_failed": False, "passed_after_retest": False}
    result = {case.id: dict(blank) for case in cases}
    if not cases or not sprint_ids:
        return result
    case_ids = [case.id for case in cases]
    runs = db.query(TestRun).filter(TestRun.sprint_id.in_(sprint_ids)).all()
    run_ids = [run.id for run in runs]
    exec_ids: set[str] = set()
    fail_execs: dict[str, int] = defaultdict(int)
    if run_ids:
        executions = (
            db.query(Execution)
            .filter(Execution.test_case_id.in_(case_ids), Execution.run_id.in_(run_ids))
            .all()
        )
        for row in executions:
            exec_ids.add(row.id)
            if row.status == "FAIL" and _in_date_range(row.executed_at, date_from, date_to):
                fail_execs[row.test_case_id] += 1
    events = (
        db.query(ExecutionEvent)
        .filter(ExecutionEvent.project_id == project_id, ExecutionEvent.test_case_id.in_(case_ids))
        .order_by(ExecutionEvent.executed_at.asc(), ExecutionEvent.id.asc())
        .all()
    )
    events_by_case: dict[str, list[ExecutionEvent]] = defaultdict(list)
    for event in events:
        if event.execution_id and event.execution_id not in exec_ids:
            continue
        if not _in_date_range(event.executed_at, date_from, date_to):
            continue
        if event.test_case_id:
            events_by_case[event.test_case_id].append(event)
    defects = (
        db.query(Defect)
        .filter(Defect.project_id == project_id, Defect.test_case_id.in_(case_ids))
        .all()
    )
    defects_by_case: dict[str, list[Defect]] = defaultdict(list)
    for defect in defects:
        if defect.execution_id and exec_ids and defect.execution_id not in exec_ids:
            continue
        if defect.test_case_id:
            defects_by_case[defect.test_case_id].append(defect)
    for case in cases:
        stats = result[case.id]
        ordered = events_by_case.get(case.id, [])
        fail_count = sum(1 for event in ordered if event.status == "FAIL")
        if fail_count == 0:
            fail_count = fail_execs.get(case.id, 0)
        if fail_count == 0 and defects_by_case.get(case.id):
            fail_count = 1
        seen_fail = False
        retest_passes = 0
        for event in ordered:
            if event.status == "FAIL":
                seen_fail = True
            elif event.status == "PASS" and seen_fail:
                retest_passes += 1
        if retest_passes == 0 and fail_count:
            current = current_by_case.get(case.id)
            current_status = current.status if current else ""
            has_pass_event = any(event.status == "PASS" for event in ordered)
            if has_pass_event:
                for defect in defects_by_case.get(case.id, []):
                    if defect_stage(defect, current_status) == "PASSED_AFTER_RETEST":
                        retest_passes = 1
                        break
        stats["fail_count"] = fail_count
        stats["retest_pass_count"] = retest_passes
        stats["ever_failed"] = fail_count > 0
        stats["passed_after_retest"] = retest_passes > 0
    return result


def folder_execute_stats(db: Session, folder: CaseFolder, run: TestRun | None = None) -> dict:
    """Unique-case current KPIs plus distinct Ever Failed / Passed After Retest for an execute folder."""
    if folder.kind == "SPRINT" and folder.sprint_id:
        cases = sprint_cases(db, folder.project_id, folder.sprint_id)
    else:
        cases = (
            db.query(TestCase)
            .filter(
                TestCase.project_id == folder.project_id,
                TestCase.folder_id == folder.id,
                TestCase.status != "DEPRECATED",
            )
            .all()
        )
    sprint_id = folder.sprint_id
    sprint = db.get(Sprint, sprint_id) if sprint_id else None
    if run is None and sprint_id:
        run = latest_run_for_sprint(db, sprint_id)
    exec_by_case: dict[str, Execution] = {}
    run_ids: list[str] = []
    if sprint_id:
        run_ids = [row.id for row in db.query(TestRun).filter(TestRun.sprint_id == sprint_id).all()]
    elif run:
        run_ids = [run.id]
    if run_ids and cases:
        rows = (
            db.query(Execution)
            .filter(Execution.run_id.in_(run_ids), Execution.test_case_id.in_([case.id for case in cases]))
            .all()
        )
        exec_by_case = _current_executions_by_case(rows)
    kpis = _kpis_for_cases(cases, exec_by_case, None, None)
    exec_counts = {
        "total": kpis["total"],
        "pass": kpis["pass"],
        "fail": kpis["fail"],
        "blocked": kpis["blocked"],
        "skipped": kpis["skipped"],
        "not_executed": kpis["not_executed"],
        "executed": kpis["executed"],
    }
    flow = empty_flow()
    if cases and sprint_id:
        history = case_history_stats(db, folder.project_id, cases, [sprint_id], None, None, exec_by_case)
        ever = sum(1 for row in history.values() if row["ever_failed"])
        after = sum(1 for row in history.values() if row["passed_after_retest"])
        flow["ever_failed"] = ever
        flow["failed"] = ever
        flow["passed_after_retest"] = after
        flow["retest_pass"] = after
        defects = (
            db.query(Defect)
            .filter(Defect.project_id == folder.project_id, Defect.test_case_id.in_([case.id for case in cases]))
            .all()
        )
        awaiting: set[str] = set()
        for defect in defects:
            if (defect.bug_id or "").strip():
                flow["bugs_raised"] += 1
            if defect_is_open(defect):
                flow["open"] += 1
            if defect.status == "CLOSED":
                flow["resolved"] += 1
            current = exec_by_case.get(defect.test_case_id or "")
            stage = defect_stage(defect, current.status if current else "")
            if defect.test_case_id and stage == "AWAITING_RETEST":
                awaiting.add(defect.test_case_id)
        flow["retest_ready"] = len(awaiting)
    return {
        "exec": exec_counts,
        "flow": flow,
        "sprint_code": sprint.code if sprint else None,
        "sprint_name": sprint.name if sprint else None,
    }


def sprint_delete_blockers(links: dict) -> str | None:
    parts = []
    if links.get("cases"):
        parts.append(f"{links['cases']} test case(s)")
    if links.get("plans"):
        parts.append(f"{links['plans']} test plan(s)")
    if links.get("runs"):
        parts.append(f"{links['runs']} execution run(s)")
    if not parts:
        return None
    return (
        "Cannot delete this sprint while it has "
        + ", ".join(parts)
        + ". Move or delete those first so execution history is kept."
    )


def _enrich_module_row(row: dict) -> dict:
    executed = row["pass"] + row["fail"] + row["blocked"] + row["skipped"]
    row["executed"] = executed
    row["pass_rate"] = round((row["pass"] / row["total"]) * 100, 1) if row["total"] else 0.0
    return row


def verdict_pass_rate(pass_n: int, fail_n: int) -> float | None:
    """Pass / (Pass + Fail). None when there are no judged results."""
    judged = pass_n + fail_n
    if not judged:
        return None
    return round((pass_n / judged) * 100, 1)


def requirement_latest_result(statuses: list[str]) -> str:
    """Roll up unique-case latest results. Fail outranks pass; mapping-only stays Not Executed."""
    if any(status == "FAIL" for status in statuses):
        return "FAIL"
    if any(status == "BLOCKED" for status in statuses):
        return "BLOCKED"
    if any(status == "PASS" for status in statuses):
        return "PASS"
    if any(status == "SKIPPED" for status in statuses):
        return "SKIPPED"
    return "NOT_EXECUTED"


def _empty_module_bucket(folder_id: str | None, name: str) -> dict:
    return {
        "id": folder_id,
        "name": name,
        "pass": 0,
        "fail": 0,
        "blocked": 0,
        "skipped": 0,
        "not_executed": 0,
        "total": 0,
        "executed": 0,
        "pass_rate": None,
    }


def _module_execution_rows(
    cases: list[TestCase],
    exec_by_case: dict[str, Execution],
    work_folders: list[CaseFolder],
    folder_id: str | None = None,
) -> list[dict]:
    module_stats: dict[str, dict] = {}
    source = work_folders if not folder_id else [f for f in work_folders if f.id == folder_id]
    for folder in source:
        module_stats[folder.id] = _empty_module_bucket(folder.id, folder.name)
    names = {f.id: f.name for f in work_folders}
    for case in cases:
        key = case.folder_id or "unassigned"
        if key not in module_stats:
            module_stats[key] = _empty_module_bucket(case.folder_id, names.get(case.folder_id or "", "Unassigned"))
        status = _latest_status(exec_by_case.get(case.id), None, None).lower()
        bucket = module_stats[key]
        bucket["total"] += 1
        bucket[status] = bucket.get(status, 0) + 1
    rows = []
    for row in module_stats.values():
        row["executed"] = row["pass"] + row["fail"] + row["blocked"] + row["skipped"]
        row["pass_rate"] = verdict_pass_rate(row["pass"], row["fail"])
        rows.append(row)
    return rows


def _bug_flow_unique(defects: list[Defect], exec_by_case: dict[str, Execution]) -> dict:
    """Unique-case history metrics; open bugs stay defect counts."""
    flow = empty_flow()
    ever: set[str] = set()
    after: set[str] = set()
    awaiting: set[str] = set()
    for defect in defects:
        case_id = defect.test_case_id
        if case_id:
            ever.add(case_id)
        if (defect.bug_id or "").strip():
            flow["bugs_raised"] += 1
        if defect_is_open(defect):
            flow["open"] += 1
        if defect.status == "CLOSED":
            flow["resolved"] += 1
        execution = exec_by_case.get(case_id) if case_id else None
        current = execution.status if execution else ""
        stage = defect_stage(defect, current)
        if case_id and stage == "AWAITING_RETEST":
            awaiting.add(case_id)
        if case_id and stage == "PASSED_AFTER_RETEST":
            after.add(case_id)
    for case_id, execution in exec_by_case.items():
        if execution.status == "FAIL":
            ever.add(case_id)
    flow["ever_failed"] = len(ever)
    flow["failed"] = len(ever)
    flow["retest_ready"] = len(awaiting)
    flow["passed_after_retest"] = len(after)
    flow["retest_pass"] = len(after)
    return flow


def dashboard(
    db: Session,
    project: Project,
    sprint_id: str | None = None,
    folder_id: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    date_scope: str | None = None,
) -> dict:
    sprints = db.query(Sprint).filter(Sprint.project_id == project.id).order_by(Sprint.created_at.asc()).all()
    chosen = None
    if sprint_id:
        chosen = next((s for s in sprints if s.id == sprint_id), None)
    if chosen is None:
        chosen = next((s for s in sprints if s.status == "ACTIVE"), sprints[-1] if sprints else None)
    if date_scope and date_scope != "sprint":
        raise HTTPException(status_code=400, detail="Invalid date_scope. Use sprint.")
    if date_scope == "sprint" and chosen and not date_from and not date_to:
        date_from = chosen.start_date
        date_to = chosen.end_date

    work_folders = (
        db.query(CaseFolder)
        .filter(CaseFolder.sprint_id == chosen.id, CaseFolder.kind == "WORK")
        .order_by(CaseFolder.created_at.asc())
        .all()
        if chosen
        else []
    )
    sprint_folder = (
        db.query(CaseFolder)
        .filter(CaseFolder.sprint_id == chosen.id, CaseFolder.kind == "SPRINT", CaseFolder.parent_id.is_(None))
        .first()
        if chosen
        else None
    )
    if folder_id:
        allowed = {f.id for f in work_folders}
        if folder_id not in allowed:
            raise HTTPException(status_code=400, detail="Module does not belong to the selected sprint")

    sprint_case_rows = sprint_cases(db, project.id, chosen.id) if chosen else []
    if folder_id:
        sprint_case_rows = [c for c in sprint_case_rows if c.folder_id == folder_id]

    run = latest_run_for_sprint(db, chosen.id) if chosen else None
    # Match Execute / RTM: all runs in the sprint, prefer a real result over a later Not Run.
    run_ids = [row.id for row in db.query(TestRun).filter(TestRun.sprint_id == chosen.id).all()] if chosen else []
    run_executions = (
        db.query(Execution).filter(Execution.run_id.in_(run_ids or ["_"])).all() if run_ids else []
    )
    exec_by_case = _current_executions_by_case(run_executions)
    kpis = _kpis_for_cases(sprint_case_rows, exec_by_case, date_from, date_to)
    completed = bool(kpis["total"] and kpis["executed"] == kpis["total"])

    sprint_series = []
    for sprint in sprints:
        sprint_run_ids = [row.id for row in db.query(TestRun).filter(TestRun.sprint_id == sprint.id).all()]
        rows = (
            db.query(Execution).filter(Execution.run_id.in_(sprint_run_ids or ["_"])).all()
            if sprint_run_ids
            else []
        )
        by_case = _current_executions_by_case(rows)
        cases = sprint_cases(db, project.id, sprint.id)
        if folder_id:
            name = next((f.name for f in work_folders if f.id == folder_id), None)
            if name:
                sprint_folder_ids = sprint_work_folder_ids(db, sprint.id)
                named = (
                    db.query(CaseFolder)
                    .filter(CaseFolder.id.in_(sprint_folder_ids or ["_"]), CaseFolder.name == name)
                    .first()
                )
                cases = [c for c in cases if named and c.folder_id == named.id]
        counts = _kpis_for_cases(cases, by_case, date_from, date_to)
        sprint_series.append(
            {
                "sprint_id": sprint.id,
                "name": sprint.name,
                "code": sprint.code,
                "status": sprint.status,
                **counts,
            }
        )

    folders = {f.id: f.name for f in work_folders}
    all_work = {
        f.id: f.name
        for f in db.query(CaseFolder).filter(CaseFolder.project_id == project.id, CaseFolder.kind == "WORK").all()
    }
    module_stats: dict[str, dict] = {}
    source_folders = work_folders if not folder_id else [f for f in work_folders if f.id == folder_id]
    for folder in source_folders:
        module_stats[folder.id] = {
            "id": folder.id,
            "name": folder.name,
            "pass": 0,
            "fail": 0,
            "blocked": 0,
            "skipped": 0,
            "not_executed": 0,
            "total": 0,
            "executed": 0,
            "pass_rate": 0.0,
        }
    for case in sprint_case_rows:
        key = case.folder_id or "unassigned"
        if key not in module_stats:
            module_stats[key] = {
                "id": case.folder_id,
                "name": all_work.get(case.folder_id or "", "Unassigned"),
                "pass": 0,
                "fail": 0,
                "blocked": 0,
                "skipped": 0,
                "not_executed": 0,
                "total": 0,
                "executed": 0,
                "pass_rate": 0.0,
            }
        status = _latest_status(exec_by_case.get(case.id), date_from, date_to).lower()
        bucket = module_stats[key]
        bucket["total"] += 1
        bucket[status] = bucket.get(status, 0) + 1
    module_rows = [_enrich_module_row(row) for row in module_stats.values() if row["total"] > 0]

    req_total = db.query(Requirement).filter(Requirement.project_id == project.id).count()
    filtered_ids = [c.id for c in sprint_case_rows]
    if filtered_ids:
        covered_ids = (
            db.query(RequirementTestCase.requirement_id)
            .join(Requirement, Requirement.id == RequirementTestCase.requirement_id)
            .filter(Requirement.project_id == project.id, RequirementTestCase.test_case_id.in_(filtered_ids))
            .distinct()
            .all()
        )
        rtm_covered = len(covered_ids)
    else:
        rtm_covered = 0
    rtm_pct = round((rtm_covered / req_total) * 100, 1) if req_total else 0.0

    users = {u.id: u.name for u in db.query(User).all()}
    visible_ids = {c.id for c in sprint_case_rows}
    recent_query = (
        db.query(Execution)
        .join(TestRun, TestRun.id == Execution.run_id)
        .filter(TestRun.project_id == project.id, Execution.status != "NOT_EXECUTED")
    )
    if chosen:
        recent_query = recent_query.filter(TestRun.sprint_id == chosen.id)
    if visible_ids:
        recent_query = recent_query.filter(Execution.test_case_id.in_(visible_ids))
    recent_rows = recent_query.order_by(Execution.executed_at.desc()).limit(20).all()
    recent_rows = [row for row in recent_rows if _in_date_range(row.executed_at, date_from, date_to)][:8]
    recent_cases = {c.id: c for c in db.query(TestCase).filter(TestCase.id.in_([e.test_case_id for e in recent_rows] or ["_"])).all()}
    recent_defects = (
        db.query(Defect).filter(Defect.execution_id.in_([e.id for e in recent_rows] or ["_"])).all()
        if recent_rows
        else []
    )
    defect_by_exec = {d.execution_id: d for d in recent_defects if d.execution_id}
    recent = []
    for execution in recent_rows:
        case = recent_cases.get(execution.test_case_id)
        defect = defect_by_exec.get(execution.id)
        stage = defect_stage(defect, execution.status) if defect else ""
        recent.append(
            {
                "id": execution.id,
                "run_id": execution.run_id,
                "case_id": case.case_id if case else "",
                "title": case.title if case else "",
                "status": execution.status,
                "executed_by_name": users.get(execution.executed_by or "", ""),
                "executed_at": iso(execution.executed_at),
                "module": all_work.get(case.folder_id or "", "") if case else "",
                "folder_id": case.folder_id if case else None,
                "retest": stage == "PASSED_AFTER_RETEST",
            }
        )

    scoped_by_case = {c.id: exec_by_case[c.id] for c in sprint_case_rows if c.id in exec_by_case}
    exec_by_id = {e.id: e for e in scoped_by_case.values()}
    sprint_defects = (
        db.query(Defect)
        .filter(Defect.project_id == project.id, Defect.test_case_id.in_([c.id for c in sprint_case_rows] or ["_"]))
        .all()
        if sprint_case_rows
        else []
    )
    bug_flow = _bug_flow_unique(sprint_defects, scoped_by_case)
    stories = []
    story_cases = {c.id: c for c in sprint_case_rows}
    for defect in sprint_defects:
        execution = scoped_by_case.get(defect.test_case_id or "") or exec_by_id.get(defect.execution_id or "")
        current = execution.status if execution else None
        case = story_cases.get(defect.test_case_id or "")
        stories.append(
            {
                "id": defect.id,
                "case_id": case.case_id if case else "",
                "title": case.title if case else defect.title,
                "bug_id": defect.bug_id or "",
                "status": defect.status,
                "original_result": defect.original_result or "FAIL",
                "current_result": current or "",
                "story": defect_story(defect, current),
            }
        )

    return {
        "project": {"id": project.id, "name": project.name, "key": project.key},
        "sprint": (
            {
                "id": chosen.id,
                "name": chosen.name,
                "code": chosen.code,
                "status": chosen.status,
                "start_date": iso(chosen.start_date),
                "end_date": iso(chosen.end_date),
            }
            if chosen
            else None
        ),
        "sprint_folder_id": sprint_folder.id if sprint_folder else None,
        "run": {"id": run.id, "name": run.name} if run else None,
        "completed": completed,
        "kpis": kpis,
        "sprint_series": sprint_series,
        "module_stats": module_rows,
        "modules": [{"id": f.id, "name": f.name} for f in work_folders],
        "rtm": {"total": req_total, "covered": rtm_covered, "percent": rtm_pct},
        "open_defects": bug_flow.get("open") or 0,
        "bug_flow": bug_flow,
        "bug_stories": stories,
        "recent": recent,
        "sprints": [
            {
                "id": s.id,
                "name": s.name,
                "code": s.code,
                "status": s.status,
                "start_date": iso(s.start_date),
                "end_date": iso(s.end_date),
            }
            for s in sprints
        ],
    }


def name_executions(db: Session, defects: list[Defect]) -> dict[str, Execution]:
    ids = {d.execution_id for d in defects if d.execution_id}
    if not ids:
        return {}
    return {row.id: row for row in db.query(Execution).filter(Execution.id.in_(ids)).all()}


def rtm_matrix(
    db: Session,
    project: Project,
    sprint_id: str | None = None,
    folder_id: str | None = None,
) -> dict:
    sprints = db.query(Sprint).filter(Sprint.project_id == project.id).order_by(Sprint.created_at.asc()).all()
    chosen = None
    if sprint_id:
        chosen = next((s for s in sprints if s.id == sprint_id), None)
    if chosen is None:
        chosen = next((s for s in sprints if s.status == "ACTIVE"), sprints[-1] if sprints else None)

    work_folders = (
        db.query(CaseFolder)
        .filter(CaseFolder.sprint_id == chosen.id, CaseFolder.kind == "WORK")
        .order_by(CaseFolder.created_at.asc())
        .all()
        if chosen
        else []
    )
    if folder_id:
        allowed = {f.id for f in work_folders}
        if folder_id not in allowed:
            raise HTTPException(status_code=400, detail="Module does not belong to the selected sprint")

    sprint_case_rows = sprint_cases(db, project.id, chosen.id) if chosen else []
    if folder_id:
        sprint_case_rows = [c for c in sprint_case_rows if c.folder_id == folder_id]

    run = latest_run_for_sprint(db, chosen.id) if chosen else None
    run_ids = [row.id for row in db.query(TestRun).filter(TestRun.sprint_id == chosen.id).all()] if chosen else []
    run_executions = (
        db.query(Execution).filter(Execution.run_id.in_(run_ids or ["_"])).all() if run_ids else []
    )
    # Same current-result rule as Execute: prefer an actual result over a later Not Run row.
    exec_by_case = _current_executions_by_case(run_executions)
    kpis = _kpis_for_cases(sprint_case_rows, exec_by_case, None, None)
    kpis["pass_rate"] = verdict_pass_rate(kpis["pass"], kpis["fail"])
    history = case_history_stats(
        db,
        project.id,
        sprint_case_rows,
        [chosen.id] if chosen else [],
        None,
        None,
        exec_by_case,
    )
    kpis["ever_failed"] = sum(1 for case in sprint_case_rows if history.get(case.id, {}).get("ever_failed"))
    kpis["passed_after_retest"] = sum(
        1 for case in sprint_case_rows if history.get(case.id, {}).get("passed_after_retest")
    )
    module_rows = _module_execution_rows(sprint_case_rows, exec_by_case, work_folders, folder_id)
    cases_by_folder: dict[str, list[TestCase]] = defaultdict(list)
    for case in sprint_case_rows:
        cases_by_folder[case.folder_id or "unassigned"].append(case)
    for row in module_rows:
        folder_cases = cases_by_folder.get(row.get("id") or "unassigned", [])
        row["ever_failed"] = sum(1 for case in folder_cases if history.get(case.id, {}).get("ever_failed"))
        row["passed_after_retest"] = sum(
            1 for case in folder_cases if history.get(case.id, {}).get("passed_after_retest")
        )
    folder_names = {
        f.id: f.name
        for f in db.query(CaseFolder).filter(CaseFolder.project_id == project.id, CaseFolder.kind == "WORK").all()
    }

    requirements = db.query(Requirement).filter(Requirement.project_id == project.id).order_by(Requirement.req_id.asc()).all()
    links = (
        db.query(RequirementTestCase)
        .join(Requirement, Requirement.id == RequirementTestCase.requirement_id)
        .filter(Requirement.project_id == project.id)
        .all()
    )
    case_ids = {link.test_case_id for link in links}
    cases = {c.id: c for c in db.query(TestCase).filter(TestCase.id.in_(case_ids or ["_"])).all()}

    links_by_req: dict[str, list[str]] = defaultdict(list)
    for link in links:
        links_by_req[link.requirement_id].append(link.test_case_id)

    rows = []
    covered = 0
    for req in requirements:
        mapped = []
        seen: set[str] = set()
        for cid in links_by_req.get(req.id, []):
            if cid in seen:
                continue
            seen.add(cid)
            case = cases.get(cid)
            if not case:
                continue
            if folder_id and case.folder_id != folder_id:
                continue
            execution = exec_by_case.get(cid)
            status = _latest_status(execution, None, None)
            mapped.append(
                {
                    "id": case.id,
                    "case_id": case.case_id,
                    "title": case.title,
                    "status": status,
                    "in_run": bool(execution),
                    "folder_id": case.folder_id,
                    "module": folder_names.get(case.folder_id or "", ""),
                }
            )
        if folder_id and not mapped:
            continue
        if mapped:
            covered += 1
        module_names: list[str] = []
        for item in mapped:
            name = item.get("module") or ""
            if name and name not in module_names:
                module_names.append(name)
        rows.append(
            {
                "id": req.id,
                "req_id": req.req_id,
                "title": req.title,
                "priority": req.priority,
                "module": ", ".join(module_names),
                "covered": bool(mapped),
                "latest_result": requirement_latest_result([item["status"] for item in mapped]),
                "cases": mapped,
            }
        )

    total = len(rows)
    return {
        "sprint_id": chosen.id if chosen else sprint_id,
        "sprint": (
            {"id": chosen.id, "name": chosen.name, "code": chosen.code, "status": chosen.status}
            if chosen
            else None
        ),
        "sprints": [{"id": s.id, "name": s.name, "code": s.code, "status": s.status} for s in sprints],
        "modules": [{"id": f.id, "name": f.name} for f in work_folders],
        "folder_id": folder_id,
        "run_id": run.id if run else None,
        "run_name": run.name if run else None,
        "kpis": kpis,
        "module_stats": module_rows,
        "total": total,
        "covered": covered,
        "percent": round((covered / total) * 100, 1) if total else 0.0,
        "rows": rows,
    }


def sprint_report(db: Session, project: Project, sprint: Sprint) -> dict:
    run = latest_run_for_sprint(db, sprint.id)
    executions = db.query(Execution).filter(Execution.run_id == run.id).all() if run else []
    cases = {c.id: c for c in db.query(TestCase).filter(TestCase.id.in_([e.test_case_id for e in executions] or ["_"])).all()}
    modules = {m.id: m.name for m in db.query(Module).filter(Module.project_id == project.id).all()}
    users = {u.id: u.name for u in db.query(User).all()}
    kpis = _counts([e.status for e in executions])

    module_stats: dict[str, dict] = defaultdict(lambda: {"name": "Unassigned", "pass": 0, "fail": 0, "blocked": 0, "skipped": 0, "not_executed": 0, "total": 0})
    tester_stats: dict[str, dict] = defaultdict(lambda: {"name": "Unassigned", "pass": 0, "fail": 0, "blocked": 0, "skipped": 0, "total": 0})
    details = []
    for execution in executions:
        case = cases.get(execution.test_case_id)
        module_name = modules.get(case.module_id, "Unassigned") if case and case.module_id else "Unassigned"
        tester = users.get(execution.executed_by or "", "Unassigned")
        module_stats[module_name]["name"] = module_name
        module_stats[module_name]["total"] += 1
        module_stats[module_name][execution.status.lower()] += 1
        if execution.status != "NOT_EXECUTED":
            tester_stats[tester]["name"] = tester
            tester_stats[tester]["total"] += 1
            tester_stats[tester][execution.status.lower()] += 1
        details.append(
            {
                "case_id": case.case_id if case else "",
                "title": case.title if case else "",
                "module": module_name,
                "priority": case.priority if case else "",
                "status": execution.status,
                "actual_result": execution.actual_result,
                "executed_by": tester if execution.executed_by else "",
                "executed_at": iso(execution.executed_at),
            }
        )

    defects = db.query(Defect).filter(Defect.project_id == project.id).all()
    defect_rows = []
    for defect in defects:
        if run and defect.execution_id and defect.execution_id not in {e.id for e in executions}:
            continue
        case = cases.get(defect.test_case_id or "")
        defect_rows.append(
            {
                "title": defect.title,
                "severity": defect.severity,
                "status": defect.status,
                "case_id": case.case_id if case else "",
            }
        )

    return {
        "project": {"name": project.name, "key": project.key},
        "sprint": {
            "id": sprint.id,
            "name": sprint.name,
            "code": sprint.code,
            "status": sprint.status,
            "start_date": iso(sprint.start_date),
            "end_date": iso(sprint.end_date),
        },
        "run": {"id": run.id, "name": run.name, "started_at": iso(run.started_at)} if run else None,
        "kpis": kpis,
        "module_stats": list(module_stats.values()),
        "tester_stats": list(tester_stats.values()),
        "details": details,
        "defects": defect_rows,
        "rtm": rtm_matrix(db, project, sprint.id),
    }


def report_excel(payload: dict) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)
    summary["A1"] = f"{payload['project']['name']} — {payload['sprint']['name']}"
    summary["A1"].font = Font(bold=True, size=14)
    kpis = payload["kpis"]
    labels = ["Total", "Executed", "Pass", "Fail", "Blocked", "Skipped", "Not executed", "Pass rate %"]
    values = [
        kpis["total"],
        kpis["executed"],
        kpis["pass"],
        kpis["fail"],
        kpis["blocked"],
        kpis["skipped"],
        kpis["not_executed"],
        kpis["pass_rate"],
    ]
    for i, (label, value) in enumerate(zip(labels, values), 3):
        summary.cell(i, 1, label)
        summary.cell(i, 2, value)

    modules = wb.create_sheet("Module-wise")
    tester = wb.create_sheet("Tester-wise")
    details = wb.create_sheet("Executions")
    for sheet, rows, cols in (
        (modules, payload["module_stats"], ["name", "total", "pass", "fail", "blocked", "skipped", "not_executed"]),
        (tester, payload["tester_stats"], ["name", "total", "pass", "fail", "blocked", "skipped"]),
        (
            details,
            payload["details"],
            ["case_id", "title", "module", "priority", "status", "actual_result", "executed_by", "executed_at"],
        ),
    ):
        for col, name in enumerate(cols, 1):
            cell = sheet.cell(1, col, name.replace("_", " ").title())
            cell.fill = header_fill
            cell.font = header_font
        for r, row in enumerate(rows, 2):
            for c, name in enumerate(cols, 1):
                sheet.cell(r, c, row.get(name, ""))

    from io import BytesIO

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _validate_report_filter(value: str | None, allowed: tuple[str, ...] | set[str], label: str) -> str | None:
    if not value:
        return None
    upper = value.upper()
    if upper not in {item.upper() for item in allowed}:
        raise HTTPException(status_code=400, detail=f"Invalid {label}")
    return upper


def execution_report(
    db: Session,
    project: Project,
    sprint_id: str | None = None,
    folder_id: str | None = None,
    case_type: str | None = None,
    priority: str | None = None,
    case_status: str | None = None,
    result: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> dict:
    case_type = _validate_report_filter(case_type, CASE_TYPES, "test type")
    priority = _validate_report_filter(priority, PRIORITIES, "priority")
    case_status = _validate_report_filter(case_status, CASE_STATUSES, "status")
    result = _validate_report_filter(result, EXEC_STATUSES, "result")

    sprints = db.query(Sprint).filter(Sprint.project_id == project.id).order_by(Sprint.created_at.desc()).all()
    selected = [sprint for sprint in sprints if sprint.id == sprint_id] if sprint_id else list(sprints)
    if sprint_id and not selected:
        raise HTTPException(status_code=404, detail="Sprint not found")
    sprint_ids = [sprint.id for sprint in selected]

    all_work = (
        db.query(CaseFolder)
        .filter(CaseFolder.sprint_id.in_([s.id for s in sprints] or ["_"]), CaseFolder.kind == "WORK")
        .order_by(CaseFolder.created_at.asc())
        .all()
        if sprints
        else []
    )
    work_folders = [folder for folder in all_work if folder.sprint_id in set(sprint_ids)]
    module_options = work_folders
    if folder_id:
        work_folders = [folder for folder in work_folders if folder.id == folder_id]
        if folder_id and not work_folders:
            raise HTTPException(status_code=400, detail="Module does not belong to the selected sprint")
    folder_ids = [folder.id for folder in work_folders]
    folder_sprint = {folder.id: folder.sprint_id for folder in work_folders}
    folder_name = {folder.id: folder.name for folder in work_folders}
    sprint_by_id = {sprint.id: sprint for sprint in sprints}

    cases_query = db.query(TestCase).filter(TestCase.project_id == project.id)
    if folder_ids:
        cases_query = cases_query.filter(TestCase.folder_id.in_(folder_ids))
    else:
        cases_query = cases_query.filter(TestCase.id == "_")
    if case_status:
        cases_query = cases_query.filter(TestCase.status == case_status)
    else:
        cases_query = cases_query.filter(TestCase.status != "DEPRECATED")
    if case_type:
        cases_query = cases_query.filter(TestCase.type == case_type)
    if priority:
        cases_query = cases_query.filter(TestCase.priority == priority)
    cases = cases_query.order_by(TestCase.case_id.asc()).all()

    all_runs = (
        db.query(TestRun)
        .filter(TestRun.sprint_id.in_(sprint_ids or ["_"]))
        .order_by(TestRun.started_at.desc())
        .all()
        if sprint_ids
        else []
    )
    latest_run: dict[str, TestRun] = {}
    for run in all_runs:
        latest_run.setdefault(run.sprint_id, run)
    all_run_ids = [run.id for run in all_runs]
    current_rows = db.query(Execution).filter(Execution.run_id.in_(all_run_ids or ["_"])).all() if all_run_ids else []
    execs_by_sprint: dict[str, list[Execution]] = defaultdict(list)
    run_sprint = {run.id: run.sprint_id for run in all_runs}
    for row in current_rows:
        sid = run_sprint.get(row.run_id)
        if sid:
            execs_by_sprint[sid].append(row)
    current_by_case: dict[str, Execution] = {}
    for case in cases:
        sid = folder_sprint.get(case.folder_id or "")
        if not sid:
            continue
        mapped = _current_executions_by_case(execs_by_sprint.get(sid, []))
        if case.id in mapped:
            current_by_case[case.id] = mapped[case.id]

    scoped_cases: list[tuple[TestCase, str]] = []
    for case in cases:
        status = _latest_status(current_by_case.get(case.id), date_from, date_to)
        if result and status != result:
            continue
        scoped_cases.append((case, status))

    kpis = _counts([status for _case, status in scoped_cases])
    kpis["pass_rate"] = verdict_pass_rate(kpis["pass"], kpis["fail"])
    history = case_history_stats(
        db,
        project.id,
        [case for case, _status in scoped_cases],
        sprint_ids,
        date_from,
        date_to,
        current_by_case,
    )
    ever_failed = sum(1 for case, _status in scoped_cases if history.get(case.id, {}).get("ever_failed"))
    passed_after = sum(1 for case, _status in scoped_cases if history.get(case.id, {}).get("passed_after_retest"))
    never_failed = max(kpis["total"] - ever_failed, 0)
    total = kpis["total"]

    users = {user.id: user.name for user in db.query(User).all()}
    last_executed_at = None
    details = []
    for case, status in scoped_cases:
        execution = current_by_case.get(case.id)
        executed_at = execution.executed_at if execution and status != "NOT_EXECUTED" else None
        if executed_at and (last_executed_at is None or executed_at > last_executed_at):
            last_executed_at = executed_at
        stats = history.get(case.id) or {}
        sid = folder_sprint.get(case.folder_id or "")
        sprint = sprint_by_id.get(sid) if sid else None
        details.append(
            {
                "id": case.id,
                "case_id": case.case_id,
                "title": case.title,
                "sprint_id": sprint.id if sprint else None,
                "sprint_code": sprint.code if sprint else "",
                "sprint_name": sprint.name if sprint else "",
                "folder_id": case.folder_id,
                "module": folder_name.get(case.folder_id or "", ""),
                "priority": case.priority,
                "type": case.type,
                "status": status,
                "ever_failed": int(stats.get("fail_count") or 0),
                "passed_after_retest": int(stats.get("retest_pass_count") or 0),
                "ever_failed_flag": bool(stats.get("ever_failed")),
                "passed_after_retest_flag": bool(stats.get("passed_after_retest")),
                "executed_at": iso(executed_at),
                "executed_by": users.get(execution.executed_by or "", "") if execution and execution.executed_by else "",
            }
        )

    cases_by_sprint: dict[str, list[tuple[TestCase, str]]] = defaultdict(list)
    for case, status in scoped_cases:
        sid = folder_sprint.get(case.folder_id or "")
        if sid:
            cases_by_sprint[sid].append((case, status))
    module_count_by_sprint: dict[str, int] = defaultdict(int)
    for folder in work_folders:
        module_count_by_sprint[folder.sprint_id] += 1

    sprint_rows = []
    for sprint in selected:
        rows = cases_by_sprint.get(sprint.id, [])
        counts = _counts([status for _case, status in rows])
        rate = verdict_pass_rate(counts["pass"], counts["fail"])
        sprint_ever = sum(1 for case, _status in rows if history.get(case.id, {}).get("ever_failed"))
        sprint_after = sum(1 for case, _status in rows if history.get(case.id, {}).get("passed_after_retest"))
        sprint_rows.append(
            {
                "id": sprint.id,
                "code": sprint.code,
                "name": sprint.name,
                "start_date": iso(sprint.start_date),
                "end_date": iso(sprint.end_date),
                "status": sprint.status,
                "module_count": module_count_by_sprint.get(sprint.id, 0),
                "total": counts["total"],
                "executed": counts["executed"],
                "pass": counts["pass"],
                "fail": counts["fail"],
                "blocked": counts["blocked"],
                "skipped": counts["skipped"],
                "not_executed": counts["not_executed"],
                "pass_rate": rate,
                "ever_failed": sprint_ever,
                "passed_after_retest": sprint_after,
            }
        )

    cases_by_folder: dict[str, list[tuple[TestCase, str]]] = defaultdict(list)
    for case, status in scoped_cases:
        if case.folder_id:
            cases_by_folder[case.folder_id].append((case, status))
    module_summary = []
    for folder in work_folders:
        rows = cases_by_folder.get(folder.id, [])
        counts = _counts([status for _case, status in rows])
        sprint = sprint_by_id.get(folder.sprint_id)
        module_summary.append(
            {
                "id": folder.id,
                "name": folder.name,
                "sprint_id": folder.sprint_id,
                "sprint_code": sprint.code if sprint else "",
                "total": counts["total"],
                "executed": counts["executed"],
                "pass": counts["pass"],
                "fail": counts["fail"],
                "blocked": counts["blocked"],
                "skipped": counts["skipped"],
                "not_executed": counts["not_executed"],
                "pass_rate": verdict_pass_rate(counts["pass"], counts["fail"]),
                "ever_failed": sum(1 for case, _status in rows if history.get(case.id, {}).get("ever_failed")),
                "passed_after_retest": sum(
                    1 for case, _status in rows if history.get(case.id, {}).get("passed_after_retest")
                ),
            }
        )

    trend_source = list(reversed(all_runs[:10])) if all_runs else []
    if date_from or date_to:
        trend_source = [
            run
            for run in reversed(all_runs)
            if _in_date_range(run.started_at, date_from, date_to)
        ][-10:]
    trend_run_ids = [run.id for run in trend_source]
    trend_execs: dict[str, list[Execution]] = defaultdict(list)
    if trend_run_ids:
        for row in db.query(Execution).filter(Execution.run_id.in_(trend_run_ids)).all():
            trend_execs[row.run_id].append(row)
    trend = []
    for run in trend_source:
        run_cases = [case for case, _status in cases_by_sprint.get(run.sprint_id, [])]
        by_case = _executions_by_case(trend_execs.get(run.id, []))
        counts = _kpis_for_cases(run_cases, by_case, date_from, date_to)
        trend.append(
            {
                "run_id": run.id,
                "name": run.name,
                "sprint_id": run.sprint_id,
                "started_at": iso(run.started_at),
                "pass": counts["pass"],
                "fail": counts["fail"],
                "blocked": counts["blocked"],
                "not_executed": counts["not_executed"],
            }
        )

    sprint_codes = {sprint.id: sprint.code for sprint in sprints}
    return {
        "project": {"id": project.id, "name": project.name, "key": project.key},
        "generated_at": iso(datetime.utcnow()),
        "last_executed_at": iso(last_executed_at),
        "filters": {
            "sprint_id": sprint_id,
            "folder_id": folder_id,
            "case_type": case_type,
            "priority": priority,
            "case_status": case_status,
            "result": result,
            "date_from": iso(date_from),
            "date_to": iso(date_to),
        },
        "sprints": [
            {
                "id": sprint.id,
                "name": sprint.name,
                "code": sprint.code,
                "status": sprint.status,
                "start_date": iso(sprint.start_date),
                "end_date": iso(sprint.end_date),
            }
            for sprint in sprints
        ],
        "modules": [
            {
                "id": folder.id,
                "name": folder.name,
                "sprint_id": folder.sprint_id,
                "sprint_code": sprint_codes.get(folder.sprint_id or "", ""),
            }
            for folder in module_options
        ],
        "kpis": {
            **kpis,
            "ever_failed": ever_failed,
            "passed_after_retest": passed_after,
            "never_failed": never_failed,
            "pass_pct": _pct_of(kpis["pass"], total),
            "fail_pct": _pct_of(kpis["fail"], total),
            "blocked_pct": _pct_of(kpis["blocked"], total),
            "not_executed_pct": _pct_of(kpis["not_executed"], total),
            "ever_failed_pct": _pct_of(ever_failed, total),
            "passed_after_retest_pct": _pct_of(passed_after, total),
            "never_failed_pct": _pct_of(never_failed, total),
        },
        "distribution": {
            "pass": kpis["pass"],
            "fail": kpis["fail"],
            "blocked": kpis["blocked"],
            "skipped": kpis["skipped"],
            "not_executed": kpis["not_executed"],
        },
        "history_summary": {
            "never_failed": never_failed,
            "ever_failed": ever_failed,
            "passed_after_retest": passed_after,
            "still_failed": max(ever_failed - passed_after, 0),
        },
        "trend": trend,
        "sprint_summary": sprint_rows,
        "module_summary": module_summary,
        "details": details,
    }


def execution_report_excel(payload: dict) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "KPIs"
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)
    project = payload["project"]
    summary["A1"] = f"{project['name']} — Execution Report"
    summary["A1"].font = Font(bold=True, size=14)
    kpis = payload["kpis"]
    labels = [
        ("Total Test Cases", kpis["total"]),
        ("Executed (Pass+Fail+Blocked+Skipped)", kpis.get("executed") or (kpis["pass"] + kpis["fail"] + kpis["blocked"] + kpis.get("skipped", 0))),
        ("Passed (Current)", kpis["pass"]),
        ("Failed (Current)", kpis["fail"]),
        ("Blocked (Current)", kpis["blocked"]),
        ("Skipped", kpis.get("skipped") or 0),
        ("Not Executed", kpis["not_executed"]),
        ("Ever Failed", kpis["ever_failed"]),
        ("Passed After Retest", kpis["passed_after_retest"]),
        ("Pass / (Pass + Fail) %", kpis["pass_rate"] if kpis["pass_rate"] is not None else "N/A"),
    ]
    for index, (label, value) in enumerate(labels, 3):
        summary.cell(index, 1, label)
        summary.cell(index, 2, value)

    sprint_sheet = wb.create_sheet("Sprint Summary")
    sprint_cols = [
        "code",
        "name",
        "start_date",
        "end_date",
        "status",
        "module_count",
        "total",
        "executed",
        "pass",
        "fail",
        "blocked",
        "skipped",
        "not_executed",
        "pass_rate",
        "ever_failed",
        "passed_after_retest",
    ]
    module_sheet = wb.create_sheet("Module Summary")
    module_cols = [
        "name",
        "sprint_code",
        "total",
        "executed",
        "pass",
        "fail",
        "blocked",
        "skipped",
        "not_executed",
        "pass_rate",
        "ever_failed",
        "passed_after_retest",
    ]
    detail_sheet = wb.create_sheet("Test Cases")
    detail_cols = [
        "case_id",
        "title",
        "sprint_code",
        "module",
        "priority",
        "type",
        "status",
        "ever_failed",
        "passed_after_retest",
        "executed_at",
    ]
    for sheet, rows, cols in (
        (sprint_sheet, payload["sprint_summary"], sprint_cols),
        (module_sheet, payload.get("module_summary") or [], module_cols),
        (detail_sheet, payload["details"], detail_cols),
    ):
        for col, name in enumerate(cols, 1):
            cell = sheet.cell(1, col, name.replace("_", " ").title())
            cell.fill = header_fill
            cell.font = header_font
        for row_idx, row in enumerate(rows, 2):
            for col, name in enumerate(cols, 1):
                value = row.get(name, "")
                if name == "pass_rate" and value is None:
                    value = "N/A"
                sheet.cell(row_idx, col, value)

    from io import BytesIO

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def execution_report_html(payload: dict) -> str:
    import html as html_lib

    project = payload["project"]
    kpis = payload["kpis"]
    pass_rate = "N/A" if kpis.get("pass_rate") is None else f"{kpis['pass_rate']}%"
    sprint_rows = "".join(
        "<tr>"
        f"<td>{html_lib.escape(str(row.get('code') or ''))}</td>"
        f"<td>{html_lib.escape(str(row.get('name') or ''))}</td>"
        f"<td>{html_lib.escape(str(row.get('status') or ''))}</td>"
        f"<td>{row.get('total') or 0}</td>"
        f"<td>{row.get('pass') or 0}</td>"
        f"<td>{row.get('fail') or 0}</td>"
        f"<td>{row.get('blocked') or 0}</td>"
        f"<td>{row.get('skipped') or 0}</td>"
        f"<td>{row.get('not_executed') or 0}</td>"
        f"<td>{'N/A' if row.get('pass_rate') is None else row.get('pass_rate')}</td>"
        f"<td>{row.get('ever_failed') or 0}</td>"
        f"<td>{row.get('passed_after_retest') or 0}</td>"
        "</tr>"
        for row in payload["sprint_summary"]
    )
    module_rows_html = "".join(
        "<tr>"
        f"<td>{html_lib.escape(str(row.get('name') or ''))}</td>"
        f"<td>{html_lib.escape(str(row.get('sprint_code') or ''))}</td>"
        f"<td>{row.get('total') or 0}</td>"
        f"<td>{row.get('pass') or 0}</td>"
        f"<td>{row.get('fail') or 0}</td>"
        f"<td>{row.get('blocked') or 0}</td>"
        f"<td>{row.get('skipped') or 0}</td>"
        f"<td>{row.get('not_executed') or 0}</td>"
        f"<td>{'N/A' if row.get('pass_rate') is None else row.get('pass_rate')}</td>"
        f"<td>{row.get('ever_failed') or 0}</td>"
        f"<td>{row.get('passed_after_retest') or 0}</td>"
        "</tr>"
        for row in payload.get("module_summary") or []
    )
    case_rows = "".join(
        "<tr>"
        f"<td>{html_lib.escape(str(row.get('case_id') or ''))}</td>"
        f"<td>{html_lib.escape(str(row.get('title') or ''))}</td>"
        f"<td>{html_lib.escape(str(row.get('sprint_code') or ''))}</td>"
        f"<td>{html_lib.escape(str(row.get('module') or ''))}</td>"
        f"<td>{html_lib.escape(str(row.get('status') or ''))}</td>"
        f"<td>{row.get('ever_failed') or 0}</td>"
        f"<td>{row.get('passed_after_retest') or 0}</td>"
        "</tr>"
        for row in payload["details"]
    )
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{html_lib.escape(project['key'])} Execution Report</title>
<style>
body {{ font-family: Inter, Segoe UI, sans-serif; color: #111827; margin: 32px; }}
h1 {{ font-size: 20px; margin: 0 0 8px; }}
h2 {{ font-size: 14px; margin: 24px 0 8px; }}
.muted {{ color: #6b7280; margin-bottom: 24px; }}
.kpis {{ display: flex; flex-wrap: wrap; gap: 12px; margin-bottom: 24px; }}
.kpi {{ border: 1px solid #e5e7eb; padding: 12px 16px; min-width: 110px; }}
.kpi b {{ display: block; font-size: 20px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 12px; }}
th, td {{ border: 1px solid #e5e7eb; padding: 6px 8px; text-align: left; }}
th {{ background: #f3f4f6; }}
@media print {{ button {{ display: none; }} }}
</style></head><body>
<button onclick="window.print()">Print / Save PDF</button>
<h1>{html_lib.escape(project['name'])} — Execution Report</h1>
<p class="muted">Unique test cases. Current result is the latest execution. Ever Failed and Passed After Retest are historical and are not added into current totals. Pass rate {pass_rate}.</p>
<div class="kpis">
<div class="kpi"><span>Total</span><b>{kpis['total']}</b></div>
<div class="kpi"><span>Executed</span><b>{kpis.get('executed') or (kpis['pass'] + kpis['fail'] + kpis['blocked'] + kpis.get('skipped', 0))}</b></div>
<div class="kpi"><span>Passed</span><b>{kpis['pass']}</b></div>
<div class="kpi"><span>Failed</span><b>{kpis['fail']}</b></div>
<div class="kpi"><span>Blocked</span><b>{kpis['blocked']}</b></div>
<div class="kpi"><span>Skipped</span><b>{kpis.get('skipped') or 0}</b></div>
<div class="kpi"><span>Not executed</span><b>{kpis['not_executed']}</b></div>
<div class="kpi"><span>Ever Failed</span><b>{kpis['ever_failed']}</b></div>
<div class="kpi"><span>Passed After Retest</span><b>{kpis['passed_after_retest']}</b></div>
</div>
<p class="muted">Executed = Pass + Fail + Blocked + Skipped (Not Executed is separate).</p>
<h2>Sprint-wise Summary</h2>
<table><thead><tr><th>Code</th><th>Name</th><th>Status</th><th>Cases</th><th>Pass</th><th>Fail</th><th>Blocked</th><th>Skipped</th><th>Not Executed</th><th>Pass Rate</th><th>Ever Failed</th><th>Passed After Retest</th></tr></thead>
<tbody>{sprint_rows or '<tr><td colspan="12">No sprints in this view.</td></tr>'}</tbody></table>
<h2>Module-wise Summary</h2>
<table><thead><tr><th>Module</th><th>Sprint</th><th>Cases</th><th>Pass</th><th>Fail</th><th>Blocked</th><th>Skipped</th><th>Not Executed</th><th>Pass Rate</th><th>Ever Failed</th><th>Passed After Retest</th></tr></thead>
<tbody>{module_rows_html or '<tr><td colspan="11">No modules in this view.</td></tr>'}</tbody></table>
<h2>Test Case Details</h2>
<table><thead><tr><th>ID</th><th>Title</th><th>Sprint</th><th>Module</th><th>Current Result</th><th>Ever Failed</th><th>Passed After Retest</th></tr></thead>
<tbody>{case_rows or '<tr><td colspan="7">No test cases in this view.</td></tr>'}</tbody></table>
</body></html>"""


def rtm_excel(payload: dict, project_name: str = "", project_key: str = "") -> bytes:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)
    summary = wb.active
    summary.title = "Summary"
    sprint = payload.get("sprint") or {}
    summary["A1"] = f"{project_name or project_key} — RTM"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"Sprint: {sprint.get('code') or '—'} {sprint.get('name') or ''}".strip()
    labels = [
        ("Requirements", payload.get("total") or 0),
        ("Covered", payload.get("covered") or 0),
        ("Coverage %", payload.get("percent") or 0),
        ("Total Cases (sprint)", (payload.get("kpis") or {}).get("total") or 0),
        ("Pass", (payload.get("kpis") or {}).get("pass") or 0),
        ("Fail", (payload.get("kpis") or {}).get("fail") or 0),
        ("Blocked", (payload.get("kpis") or {}).get("blocked") or 0),
        ("Not Executed", (payload.get("kpis") or {}).get("not_executed") or 0),
    ]
    for index, (label, value) in enumerate(labels, 4):
        summary.cell(index, 1, label)
        summary.cell(index, 2, value)

    matrix = wb.create_sheet("RTM Matrix")
    cols = [
        "Req ID",
        "Requirement",
        "Priority",
        "Module",
        "Covered",
        "Latest Result",
        "Case ID",
        "Case Title",
        "Case Result",
        "Case Module",
    ]
    for col, name in enumerate(cols, 1):
        cell = matrix.cell(1, col, name)
        cell.fill = header_fill
        cell.font = header_font
    row_idx = 2
    for req in payload.get("rows") or []:
        cases = req.get("cases") or []
        if not cases:
            matrix.cell(row_idx, 1, req.get("req_id") or "")
            matrix.cell(row_idx, 2, req.get("title") or "")
            matrix.cell(row_idx, 3, req.get("priority") or "")
            matrix.cell(row_idx, 4, req.get("module") or "")
            matrix.cell(row_idx, 5, "Yes" if req.get("covered") else "No")
            matrix.cell(row_idx, 6, req.get("latest_result") or "NOT_EXECUTED")
            row_idx += 1
            continue
        for case in cases:
            matrix.cell(row_idx, 1, req.get("req_id") or "")
            matrix.cell(row_idx, 2, req.get("title") or "")
            matrix.cell(row_idx, 3, req.get("priority") or "")
            matrix.cell(row_idx, 4, req.get("module") or "")
            matrix.cell(row_idx, 5, "Yes" if req.get("covered") else "No")
            matrix.cell(row_idx, 6, req.get("latest_result") or "NOT_EXECUTED")
            matrix.cell(row_idx, 7, case.get("case_id") or "")
            matrix.cell(row_idx, 8, case.get("title") or "")
            matrix.cell(row_idx, 9, case.get("status") or "NOT_EXECUTED")
            matrix.cell(row_idx, 10, case.get("module") or "")
            row_idx += 1

    from io import BytesIO

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rtm_html(payload: dict, project_name: str = "", project_key: str = "") -> str:
    import html as html_lib

    sprint = payload.get("sprint") or {}
    rows_html = []
    for req in payload.get("rows") or []:
        cases = req.get("cases") or []
        if not cases:
            rows_html.append(
                "<tr>"
                f"<td>{html_lib.escape(str(req.get('req_id') or ''))}</td>"
                f"<td>{html_lib.escape(str(req.get('title') or ''))}</td>"
                f"<td>{html_lib.escape(str(req.get('priority') or ''))}</td>"
                f"<td>{html_lib.escape(str(req.get('module') or ''))}</td>"
                f"<td>{'Yes' if req.get('covered') else 'No'}</td>"
                f"<td>{html_lib.escape(str(req.get('latest_result') or 'NOT_EXECUTED'))}</td>"
                "<td>—</td><td>—</td><td>—</td>"
                "</tr>"
            )
            continue
        for case in cases:
            rows_html.append(
                "<tr>"
                f"<td>{html_lib.escape(str(req.get('req_id') or ''))}</td>"
                f"<td>{html_lib.escape(str(req.get('title') or ''))}</td>"
                f"<td>{html_lib.escape(str(req.get('priority') or ''))}</td>"
                f"<td>{html_lib.escape(str(req.get('module') or ''))}</td>"
                f"<td>{'Yes' if req.get('covered') else 'No'}</td>"
                f"<td>{html_lib.escape(str(req.get('latest_result') or 'NOT_EXECUTED'))}</td>"
                f"<td>{html_lib.escape(str(case.get('case_id') or ''))}</td>"
                f"<td>{html_lib.escape(str(case.get('title') or ''))}</td>"
                f"<td>{html_lib.escape(str(case.get('status') or 'NOT_EXECUTED'))}</td>"
                "</tr>"
            )
    body_rows = "".join(rows_html)
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{html_lib.escape(project_key or 'RTM')} RTM</title>
<style>
body {{ font-family: Inter, Segoe UI, sans-serif; color: #111827; margin: 32px; }}
h1 {{ font-size: 20px; margin: 0 0 8px; }}
.muted {{ color: #6b7280; margin-bottom: 16px; }}
.kpis {{ display: flex; flex-wrap: wrap; gap: 12px; margin-bottom: 20px; }}
.kpi {{ border: 1px solid #e5e7eb; padding: 10px 14px; min-width: 100px; }}
.kpi b {{ display: block; font-size: 18px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 12px; }}
th, td {{ border: 1px solid #e5e7eb; padding: 6px 8px; text-align: left; }}
th {{ background: #f3f4f6; }}
@media print {{ button {{ display: none; }} }}
</style></head><body>
<button onclick="window.print()">Print / Save PDF</button>
<h1>{html_lib.escape(project_name or project_key)} — Requirements Traceability Matrix</h1>
<p class="muted">{html_lib.escape(str(sprint.get('code') or ''))} {html_lib.escape(str(sprint.get('name') or ''))} · req × cases × latest result</p>
<div class="kpis">
<div class="kpi"><span>Requirements</span><b>{payload.get('total') or 0}</b></div>
<div class="kpi"><span>Covered</span><b>{payload.get('covered') or 0}</b></div>
<div class="kpi"><span>Coverage</span><b>{payload.get('percent') or 0}%</b></div>
</div>
<table><thead><tr>
<th>Req ID</th><th>Requirement</th><th>Priority</th><th>Module</th><th>Covered</th><th>Latest Result</th>
<th>Case ID</th><th>Case Title</th><th>Case Result</th>
</tr></thead>
<tbody>{body_rows or '<tr><td colspan="9">No requirements in this view.</td></tr>'}</tbody></table>
</body></html>"""

