import csv
import io
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.ids import next_case_id, next_req_id
from app.models import Module, Project, Requirement, RequirementTestCase, TestCase, TestStep, User

CANONICAL = [
    "testcaseid",
    "module",
    "title",
    "description",
    "preconditions",
    "testdata",
    "priority",
    "type",
    "requirementids",
    "status",
    "steps",
    "action",
    "expectedresult",
    "stepno",
]

ALIASES = {
    "test case id": "testcaseid",
    "tc id": "testcaseid",
    "id": "testcaseid",
    "case id": "testcaseid",
    "module name": "module",
    "feature": "module",
    "summary": "title",
    "name": "title",
    "test case": "title",
    "desc": "description",
    "pre-conditions": "preconditions",
    "precondition": "preconditions",
    "pre-condition": "preconditions",
    "test data": "testdata",
    "data": "testdata",
    "priority": "priority",
    "type": "type",
    "test type": "type",
    "requirement ids": "requirementids",
    "requirement id": "requirementids",
    "requirements": "requirementids",
    "req id": "requirementids",
    "reqids": "requirementids",
    "status": "status",
    "steps": "steps",
    "step": "steps",
    "action": "action",
    "step action": "action",
    "expected result": "expectedresult",
    "expected": "expectedresult",
    "expected results": "expectedresult",
    "step no": "stepno",
    "step number": "stepno",
    "#": "stepno",
}


def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (value or "").strip().lower())


def auto_map(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for header in headers:
        key = _norm(header)
        spaced = re.sub(r"\s+", " ", (header or "").strip().lower())
        target = ALIASES.get(spaced) or (key if key in CANONICAL else None)
        if target:
            mapping[header] = target
    return mapping


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases"
    headers = [
        "TestCaseID",
        "Module",
        "Title",
        "Description",
        "Preconditions",
        "TestData",
        "Priority",
        "Type",
        "RequirementIDs",
        "Status",
        "Steps",
    ]
    fill = PatternFill("solid", fgColor="2563EB")
    font = Font(color="FFFFFF", bold=True)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 22
    ws.append(
        [
            "",
            "Login",
            "Valid user can log in",
            "User with valid credentials reaches home",
            "User exists and is active",
            "user@shop.test / Pass@123",
            "High",
            "Functional",
            "REQ-SHOP-001",
            "Ready",
            "1. Open login page || Login form is shown\n2. Enter email and password || Fields accept input\n3. Click Sign in || Home page loads",
        ]
    )
    steps = wb.create_sheet("StepsLayout")
    step_headers = ["TestCaseID", "Module", "Title", "Priority", "Type", "RequirementIDs", "StepNo", "Action", "ExpectedResult"]
    for col, header in enumerate(step_headers, 1):
        cell = steps.cell(1, col, header)
        cell.fill = fill
        cell.font = font
    steps.append(["TC-DEMO-001", "Login", "Valid user can log in", "High", "Functional", "REQ-SHOP-001", 1, "Open login page", "Login form is shown"])
    steps.append(["TC-DEMO-001", "Login", "Valid user can log in", "High", "Functional", "REQ-SHOP-001", 2, "Enter email and password", "Fields accept input"])
    help_sheet = wb.create_sheet("Help")
    help_sheet["A1"] = "Use TestCases sheet (one row per case) or StepsLayout (one row per step, same TestCaseID)."
    help_sheet["A2"] = "In one-row format, split Action and Expected Result with || inside the Steps cell."
    help_sheet["A3"] = "TestCaseID is optional — the app generates TC-{PROJECT}-{nnn} when empty."
    help_sheet["A4"] = "Priority: High / Medium / Low. Type: Functional / Regression / Smoke / Sanity / Negative / UAT."
    help_sheet["A5"] = "RequirementIDs can be comma-separated. Unknown IDs are created as stub requirements."
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(row: dict, mapping: dict[str, str], field: str) -> str:
    for header, target in mapping.items():
        if target == field:
            value = row.get(header)
            return "" if value is None else str(value).strip()
    return ""


def parse_steps(text: str) -> list[tuple[str, str]]:
    steps: list[tuple[str, str]] = []
    for raw in re.split(r"[\r\n]+", text or ""):
        line = re.sub(r"^\s*\d+[\).\:\-]\s*", "", raw).strip()
        if not line:
            continue
        if "||" in line:
            action, expected = line.split("||", 1)
            steps.append((action.strip(), expected.strip()))
        else:
            steps.append((line, ""))
    return steps


def read_tabular(filename: str, data: bytes) -> tuple[list[str], list[dict]]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        headers = reader.fieldnames or []
        rows = [dict(r) for r in reader]
        return list(headers), rows
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    rows_iter = list(ws.iter_rows(values_only=True))
    if not rows_iter:
        return [], []
    headers = [str(h).strip() if h is not None else f"Column{i}" for i, h in enumerate(rows_iter[0], 1)]
    rows = []
    for raw in rows_iter[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        rows.append({headers[i]: ("" if raw[i] is None else raw[i]) for i in range(len(headers)) if i < len(raw)})
    return headers, rows


def group_rows(rows: list[dict], mapping: dict[str, str]) -> list[dict]:
    grouped: dict[str, dict] = {}
    standalone: list[dict] = []
    for index, row in enumerate(rows):
        title = _cell(row, mapping, "title")
        action = _cell(row, mapping, "action")
        case_key = _cell(row, mapping, "testcaseid") or title or f"row-{index}"
        if action and not _cell(row, mapping, "steps"):
            bucket = grouped.setdefault(
                case_key,
                {
                    "key": case_key,
                    "title": title,
                    "module": _cell(row, mapping, "module"),
                    "description": _cell(row, mapping, "description"),
                    "preconditions": _cell(row, mapping, "preconditions"),
                    "testdata": _cell(row, mapping, "testdata"),
                    "priority": _cell(row, mapping, "priority"),
                    "type": _cell(row, mapping, "type"),
                    "requirementids": _cell(row, mapping, "requirementids"),
                    "status": _cell(row, mapping, "status"),
                    "steps": [],
                },
            )
            if title:
                bucket["title"] = title
            bucket["steps"].append(
                {
                    "step_no": int(_cell(row, mapping, "stepno") or len(bucket["steps"]) + 1),
                    "action": action,
                    "expected_result": _cell(row, mapping, "expectedresult"),
                }
            )
        else:
            parsed = parse_steps(_cell(row, mapping, "steps"))
            standalone.append(
                {
                    "key": case_key,
                    "title": title,
                    "module": _cell(row, mapping, "module"),
                    "description": _cell(row, mapping, "description"),
                    "preconditions": _cell(row, mapping, "preconditions"),
                    "testdata": _cell(row, mapping, "testdata"),
                    "priority": _cell(row, mapping, "priority"),
                    "type": _cell(row, mapping, "type"),
                    "requirementids": _cell(row, mapping, "requirementids"),
                    "status": _cell(row, mapping, "status"),
                    "steps": [
                        {"step_no": i + 1, "action": a, "expected_result": e} for i, (a, e) in enumerate(parsed)
                    ],
                }
            )
    return standalone + list(grouped.values())


def _norm_priority(value: str) -> str:
    v = (value or "MEDIUM").upper().strip()
    return v if v in {"HIGH", "MEDIUM", "LOW"} else "MEDIUM"


def _norm_type(value: str) -> str:
    v = (value or "FUNCTIONAL").upper().strip()
    allowed = {"FUNCTIONAL", "REGRESSION", "SMOKE", "SANITY", "NEGATIVE", "UAT"}
    return v if v in allowed else "FUNCTIONAL"


def _norm_status(value: str) -> str:
    v = (value or "READY").upper().strip().replace(" ", "_")
    if v == "READYFOREXECUTE":
        v = "READY_FOR_EXECUTE"
    return v if v in {"DRAFT", "READY", "READY_FOR_EXECUTE", "DEPRECATED"} else "READY"


def _get_or_create_module(db: Session, project_id: str, name: str, cache: dict[str, Module]) -> Module | None:
    clean = (name or "").strip()
    if not clean:
        return None
    key = clean.lower()
    if key in cache:
        return cache[key]
    existing = (
        db.query(Module)
        .filter(Module.project_id == project_id, Module.name == clean)
        .first()
    )
    if not existing:
        existing = Module(project_id=project_id, name=clean)
        db.add(existing)
        db.flush()
    cache[key] = existing
    return existing


def classify_import_cases(
    cases: list[dict],
    existing_ids: set[str],
    folder_name: str | None = None,
) -> dict:
    """Validate grouped import cases before commit. Status: ok | error | duplicate_file | duplicate_existing | bad_module."""
    items: list[dict] = []
    issues: list[dict] = []
    seen: dict[str, int] = {}
    existing_cf = {value.casefold() for value in existing_ids}
    folder_cf = (folder_name or "").strip().casefold() or None

    for index, case in enumerate(cases, 1):
        title = (case.get("title") or "").strip()
        module = (case.get("module") or "").strip()
        requested_id = (case.get("key") or "").strip()
        looks_like_id = bool(requested_id) and requested_id.upper().startswith("TC-")
        case_id = requested_id if looks_like_id else ""
        item = {
            "row": index,
            "case_id": case_id,
            "title": title,
            "module": module,
            "case": case,
        }
        if not title:
            item["status"] = "error"
            item["issue"] = "Title is required"
            issues.append({"row": index, "case_id": case_id, "message": "Title is required"})
        elif looks_like_id and case_id.casefold() in seen:
            item["status"] = "duplicate_file"
            item["issue"] = f"Duplicate case ID in file (also row {seen[case_id.casefold()]})"
            issues.append({"row": index, "case_id": case_id, "message": "Duplicate Case ID in this file"})
        elif looks_like_id and case_id.casefold() in existing_cf:
            item["status"] = "duplicate_existing"
            item["issue"] = "Case ID already exists"
            issues.append({"row": index, "case_id": case_id, "message": "Case ID already exists in this project"})
        elif module and folder_cf and module.casefold() != folder_cf:
            item["status"] = "bad_module"
            item["issue"] = f"Module “{module}” does not match “{folder_name}”"
            issues.append(
                {
                    "row": index,
                    "case_id": case_id,
                    "message": f"Module “{module}” does not match import folder “{folder_name}”",
                }
            )
        else:
            item["status"] = "ok"
            if looks_like_id:
                seen[case_id.casefold()] = index
        items.append(item)

    return {
        "items": items,
        "issues": issues,
        "ok_count": sum(1 for item in items if item.get("status") == "ok"),
        "error_count": sum(1 for item in items if item.get("status") != "ok"),
    }


def commit_import(
    db: Session,
    project: Project,
    user: User,
    cases: list[dict],
    folder_id: str | None = None,
    folder_name: str | None = None,
) -> dict:
    existing_ids = {
        row.case_id
        for row in db.query(TestCase).filter(TestCase.project_id == project.id).all()
        if row.case_id
    }
    classified = classify_import_cases(cases, existing_ids, folder_name=folder_name)
    module_cache: dict[str, Module] = {}
    created = 0
    skipped = 0
    for item in classified["items"]:
        if item.get("status") != "ok":
            skipped += 1
            continue
        case_data = item["case"]
        title = (case_data.get("title") or "").strip()
        module = _get_or_create_module(db, project.id, case_data.get("module") or folder_name or "", module_cache)
        requested_id = (case_data.get("key") or "").strip()
        if requested_id and requested_id.upper().startswith("TC-"):
            case_id = requested_id
        else:
            case_id = next_case_id(db, project)
        case = TestCase(
            project_id=project.id,
            module_id=module.id if module else None,
            folder_id=folder_id,
            case_id=case_id,
            title=title,
            description=case_data.get("description") or "",
            preconditions=case_data.get("preconditions") or "",
            test_data=case_data.get("testdata") or "",
            priority=_norm_priority(case_data.get("priority") or ""),
            type=_norm_type(case_data.get("type") or ""),
            status=_norm_status(case_data.get("status") or ""),
            created_by=user.id,
        )
        db.add(case)
        db.flush()
        existing_ids.add(case_id)
        for step in case_data.get("steps") or []:
            db.add(
                TestStep(
                    test_case_id=case.id,
                    step_no=int(step.get("step_no") or 1),
                    action=step.get("action") or "",
                    expected_result=step.get("expected_result") or "",
                )
            )
        req_blob = case_data.get("requirementids") or ""
        for code in [p.strip() for p in re.split(r"[,;]+", req_blob) if p.strip()]:
            req = (
                db.query(Requirement)
                .filter(Requirement.project_id == project.id, Requirement.req_id == code)
                .first()
            )
            if not req:
                if not code.upper().startswith("REQ-"):
                    code = next_req_id(db, project)
                    req = Requirement(project_id=project.id, req_id=code, title=title, description="Created from import")
                else:
                    req = Requirement(project_id=project.id, req_id=code, title=title, description="Created from import")
                db.add(req)
                db.flush()
            db.add(RequirementTestCase(requirement_id=req.id, test_case_id=case.id))
        created += 1
    db.commit()
    return {"created": created, "skipped": skipped}
